"""
Lightweight automated checks for the KPI Dashboard's Executive Dashboard
layout fix: chart anchors are unique, chart bounding boxes (computed from
actual column widths / row heights, not guessed) never overlap, and the
worksheet is configured to scroll normally rather than behave like a
frozen, wide horizontal canvas.

Run with: python -m unittest discover -s tests -v
(Requires output/kpi_dashboard.xlsx to exist — run `make all` first.)
"""

import sys
import unittest
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))

import openpyxl  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.utils.cell import coordinate_to_tuple  # noqa: E402

from close_pack.excel_export import _col_width_to_px, _row_height_to_px  # noqa: E402

EMU_PER_PX = 9525
KPI_DASHBOARD_PATH = REPO_ROOT / "output" / "kpi_dashboard.xlsx"


def _image_bbox_px(ws, img) -> tuple:
    """Absolute pixel bounding box (x0, y0, x1, y1) of an embedded image,
    from its anchor's 0-indexed (col, row) plus the EXPLICIT display size
    stored in anchor.ext (EMU) — the same geometry Excel actually renders,
    not the PNG's native pixel size (which openpyxl re-reports after a
    reload regardless of the display size set at write time)."""
    anchor_from = img.anchor._from
    x0 = sum(
        _col_width_to_px(ws.column_dimensions[get_column_letter(c + 1)].width or 8.43)
        for c in range(anchor_from.col)
    )
    y0 = sum(
        _row_height_to_px(ws.row_dimensions[r + 1].height or 15.0)
        for r in range(anchor_from.row)
    )
    width_px = img.anchor.ext.cx / EMU_PER_PX
    height_px = img.anchor.ext.cy / EMU_PER_PX
    return (x0, y0, x0 + width_px, y0 + height_px)


def _boxes_overlap(box_a: tuple, box_b: tuple) -> bool:
    ax0, ay0, ax1, ay1 = box_a
    bx0, by0, bx1, by1 = box_b
    return ax0 < bx1 and bx0 < ax1 and ay0 < by1 and by0 < ay1


class TestKPIDashboardLayout(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not KPI_DASHBOARD_PATH.exists():
            raise unittest.SkipTest(f"{KPI_DASHBOARD_PATH} not found — run `make all` first.")
        cls.wb = openpyxl.load_workbook(KPI_DASHBOARD_PATH)
        cls.ws = cls.wb["Dashboard"]

    def test_exactly_three_charts(self):
        self.assertEqual(len(self.ws._images), 3, "Executive Dashboard must use exactly 3 charts")

    def test_chart_anchors_are_unique(self):
        anchors = [(img.anchor._from.col, img.anchor._from.row) for img in self.ws._images]
        self.assertEqual(len(anchors), len(set(anchors)), f"duplicate chart anchor cells: {anchors}")

    def test_chart_bounding_boxes_do_not_overlap(self):
        boxes = [_image_bbox_px(self.ws, img) for img in self.ws._images]
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                self.assertFalse(
                    _boxes_overlap(boxes[i], boxes[j]),
                    f"chart {i} {boxes[i]} overlaps chart {j} {boxes[j]}",
                )

    def test_chart_images_fit_within_their_own_target_range(self):
        """Each image's explicit size must not itself be absurdly larger
        than a single dashboard grid cell block (a regression guard against
        reverting to native-PNG-size embedding, which is what caused the
        original overlap bug)."""
        for img in self.ws._images:
            width_px = img.anchor.ext.cx / EMU_PER_PX
            height_px = img.anchor.ext.cy / EMU_PER_PX
            self.assertLess(width_px, 1200, "chart width looks like an unsized native PNG, not an explicit size")
            self.assertLess(height_px, 700, "chart height looks like an unsized native PNG, not an explicit size")

    def test_freeze_pane_is_modest(self):
        """A large freeze (e.g. freezing the KPI table's own header row,
        deep in the sheet) is what made the dashboard feel like a static,
        non-scrolling canvas. Only a small top-row freeze is allowed."""
        freeze = self.ws.freeze_panes
        self.assertIsNotNone(freeze, "dashboard should have a simple top-row freeze")
        row, col = coordinate_to_tuple(freeze)
        self.assertLessEqual(row, 6, f"freeze pane at row {row} is not modest — it pins too much of the sheet")
        self.assertEqual(col, 1, "freeze pane should not freeze any columns")

    def test_vertical_and_horizontal_scrolling_enabled(self):
        self.assertGreaterEqual(len(self.wb.views), 1)
        view = self.wb.views[0]
        self.assertTrue(view.showVerticalScroll, "vertical scrollbar must be enabled")
        self.assertTrue(view.showHorizontalScroll, "horizontal scrollbar must be enabled")

    def test_sheet_view_is_normal(self):
        view = self.ws.sheet_view.view
        self.assertIn(view, (None, "normal"), f"sheet should render in normal view, got {view!r}")

    def test_content_extends_past_the_freeze_into_a_tall_sheet(self):
        """Confirms the dashboard is a genuinely tall, scrollable sheet
        (cards -> insights -> a 2x2 chart/table grid down to row 47), not
        content crammed into the first handful of rows."""
        max_row_with_content = max(
            (img.anchor._from.row + 1 for img in self.ws._images), default=0
        )
        self.assertGreaterEqual(max_row_with_content, 33,
                                 "dashboard content should extend well below the frozen title rows")


if __name__ == "__main__":
    unittest.main()
