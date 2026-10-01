"""
Targeted validation tests added for the financial-safeguard review:
    1. Operating Income calculation and reconciliation
    2. Favorable/unfavorable sign logic
    3. Waterfall reconciles to Operating Income variance (after rounding)
    4. No calculated subtotal is used as a waterfall driver
    5. YoY growth is not exposed anywhere without genuine 24-month history
    6. Forecast Accuracy has no look-ahead bias (forecast for a month never
       depends on that month's, or a later month's, actual data)
    7. Workbook generation completes successfully end to end
    8. Percent variance is N/M when a profit metric (Gross Profit/Operating
       Income) flips sign vs. its comparison, even if materially flagged
    9. Executive commentary (material_top_drivers / narrative) never cites
       an immaterial line item

Run with: python -m unittest discover -s tests -v
"""

import sys
import tempfile
import unittest
from pathlib import Path

import warnings
import zipfile

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, str(REPO_ROOT / "data"))

import pandas as pd  # noqa: E402

from close_pack import controls, excel_export, kpi, variance  # noqa: E402


def _real_pnl_and_variance():
    var_df = variance.compute_variance(variance.load_pnl(REPO_ROOT / "data" / "close_pack.db"))
    return var_df


class TestOperatingIncomeReconciliation(unittest.TestCase):
    def test_operating_income_equals_gross_profit_minus_opex(self):
        """OI = Gross Profit - sum(opex), for actual/budget/forecast, every month."""
        var_df = _real_pnl_and_variance()
        for month, group in var_df.groupby("month"):
            by_line = group.set_index("line_item")
            for scenario in ("actual", "budget", "forecast"):
                gross_profit = by_line.loc["Revenue", scenario] - by_line.loc["COGS", scenario]
                opex_total = by_line.loc[variance.OPEX_LINES, scenario].sum()
                expected_oi = gross_profit - opex_total
                actual_oi = by_line.loc["Operating Income", scenario]
                self.assertAlmostEqual(expected_oi, actual_oi, places=2,
                                        msg=f"OI reconciliation failed for {month}/{scenario}")


class TestFavorableUnfavorableSignLogic(unittest.TestCase):
    def setUp(self):
        rows = [
            {"month": "2099-01", "line_item": "Revenue", "actual": 110_000, "budget": 100_000, "forecast": 100_000},
            {"month": "2099-01", "line_item": "COGS", "actual": 55_000, "budget": 50_000, "forecast": 50_000},
            {"month": "2099-01", "line_item": "Salaries", "actual": 20_000, "budget": 20_000, "forecast": 20_000},
            {"month": "2099-01", "line_item": "Marketing", "actual": 20_000, "budget": 20_000, "forecast": 20_000},
            {"month": "2099-01", "line_item": "Facilities", "actual": 5_000, "budget": 5_000, "forecast": 5_000},
            {"month": "2099-01", "line_item": "Professional Services", "actual": 3_000, "budget": 3_000, "forecast": 3_000},
            {"month": "2099-01", "line_item": "Other", "actual": 1_000, "budget": 1_000, "forecast": 1_000},
        ]
        pnl = variance.add_subtotals(pd.DataFrame(rows))
        self.var_df = variance.compute_variance(pnl)

    def _favorable(self, line_item):
        row = self.var_df[self.var_df.line_item == line_item].iloc[0]
        return row["favorable_variance_vs_budget"]

    def test_higher_revenue_is_favorable(self):
        self.assertGreater(self._favorable("Revenue"), 0)

    def test_higher_cogs_is_unfavorable(self):
        self.assertLess(self._favorable("COGS"), 0)

    def test_higher_operating_income_is_favorable(self):
        # Revenue +10k, COGS +5k (cost overrun) -> GP +5k, opex unchanged -> OI +5k favorable
        self.assertGreater(self._favorable("Operating Income"), 0)

    def test_lower_opex_would_be_favorable(self):
        rows = [
            {"month": "2099-02", "line_item": "Marketing", "actual": 8_000, "budget": 10_000, "forecast": 10_000},
        ]
        df = pd.DataFrame(rows)
        result = variance.compute_variance(df)
        self.assertGreater(result.iloc[0]["favorable_variance_vs_budget"], 0,
                            "spending LESS than budget on an opex line must be favorable")


class TestWaterfallReconciliation(unittest.TestCase):
    def test_reconciles_after_rounding_for_every_month(self):
        var_df = _real_pnl_and_variance()
        for month in sorted(var_df["month"].unique()):
            recon = variance.waterfall_reconciliation(var_df, month)
            self.assertEqual(recon["difference"], 0.0,
                              msg=f"waterfall does not reconcile to $0.00 for {month}: {recon}")

    def test_no_calculated_subtotal_used_as_a_driver(self):
        """Gross Profit, Total Opex, and Operating Income must never appear
        as their own waterfall bar — they are calculated subtotals of the
        bars already shown, and including them would double-count."""
        forbidden = {"Gross Profit", "Operating Income", "Total Opex", "Total OpEx"}
        self.assertTrue(forbidden.isdisjoint(set(variance.BRIDGE_ORDER)),
                         f"BRIDGE_ORDER must contain only leaf-level lines, found: "
                         f"{forbidden.intersection(variance.BRIDGE_ORDER)}")
        # and every line actually in the bridge is a real leaf-level P&L input
        leaf_lines = {"Revenue", "COGS"} | set(variance.OPEX_LINES)
        self.assertTrue(set(variance.BRIDGE_ORDER).issubset(leaf_lines))


class TestNoUnsupportedYoY(unittest.TestCase):
    def test_yoy_not_in_kpi_registry(self):
        keys = [spec.key for spec in kpi.KPI_REGISTRY]
        self.assertNotIn("revenue_growth_yoy", keys,
                          "YoY must not be exposed as a KPI: this dataset has only one fiscal year "
                          "of real monthly actuals, not genuine 24-month history.")

    def test_kpi_module_has_no_yoy_function(self):
        self.assertFalse(hasattr(kpi, "revenue_growth_yoy"))

    def test_no_prior_year_table_in_schema(self):
        import sqlite3
        conn = sqlite3.connect(REPO_ROOT / "data" / "close_pack.db")
        try:
            tables = {row[0] for row in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table'")}
        finally:
            conn.close()
        self.assertNotIn("prior_year_revenue", tables,
                          "a synthetic prior-year table would be fabricated data dressed up as history")


class TestForecastAccuracyValidity(unittest.TestCase):
    """Forecast Accuracy is only a valid KPI if the forecast used for a
    given month was actually available BEFORE that month happened — i.e.
    no look-ahead. generate_data._generate_forecast revises the forecast
    quarterly using only the ELAPSED quarters' actuals; this reconstructs
    that computation independently and confirms no leakage."""

    def test_forecast_never_uses_same_or_later_month_actuals(self):
        import generate_data

        data = generate_data.generate()
        actual = data["actual"]
        budget = data["budget"]
        forecast = data["forecast"]
        quarter_starts = [0, 3, 6, 9]

        for line in generate_data.LINE_ITEMS:
            expected = list(budget[line])  # Q1 vintage == budget
            for start in quarter_starts[1:]:
                elapsed_actual_sum = sum(actual[line][0:start])
                elapsed_budget_sum = sum(budget[line][0:start])
                factor = elapsed_actual_sum / elapsed_budget_sum if elapsed_budget_sum else 1.0
                factor = min(max(factor, 0.85), 1.15)
                for m_idx in range(start, 12):
                    expected[m_idx] = round(budget[line][m_idx] * factor, 2)

            for m_idx in range(12):
                self.assertAlmostEqual(
                    forecast[line][m_idx], expected[m_idx], places=2,
                    msg=f"{line} month {m_idx}: forecast does not match the documented "
                        f"elapsed-quarters-only computation (possible look-ahead leak)",
                )

    def test_forecast_accuracy_kpi_is_well_defined(self):
        inp = kpi.load_kpi_inputs(REPO_ROOT / "data" / "close_pack.db")
        series = kpi.forecast_accuracy_pct(inp)
        self.assertEqual(len(series), 12)
        self.assertTrue((series >= 0).all())
        self.assertTrue((series <= 1).all())


class TestWorkbookGenerationSucceeds(unittest.TestCase):
    def test_full_pipeline_produces_valid_workbooks(self):
        import generate_data

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            db_path = tmp_path / "close_pack.db"
            data = generate_data.generate()
            generate_data._write_sqlite(data, db_path)

            var_df = variance.compute_variance(variance.load_pnl(db_path))
            kpi_inputs = kpi.load_kpi_inputs(db_path)
            calendar_df = controls.close_calendar()
            controls_df = controls.evaluate_controls(str(db_path))

            month = generate_data.MONTHS[-1]
            variance_path = tmp_path / "variance_analysis.xlsx"
            kpi_path = tmp_path / "kpi_dashboard.xlsx"
            close_path = tmp_path / "close_checklist.xlsx"

            excel_export.build_variance_workbook(var_df, month, variance_path)
            excel_export.build_kpi_workbook(kpi_inputs, month, kpi_path)
            excel_export.build_close_workbook(calendar_df, controls_df, close_path)

            for path in (variance_path, kpi_path, close_path):
                self.assertTrue(path.exists(), f"{path} was not created")
                zf = zipfile.ZipFile(path)
                self.assertIsNone(zf.testzip(), f"{path} failed zip integrity check")

                with warnings.catch_warnings():
                    warnings.simplefilter("error")
                    import openpyxl
                    wb = openpyxl.load_workbook(path)
                    self.assertGreaterEqual(len(wb.sheetnames), 1)


class TestSignFlipNotMeaningful(unittest.TestCase):
    """A profit metric (Gross Profit, Operating Income) that swings from
    budgeted profit to an actual loss (or vice versa) produces a % variance
    that is technically computable but not a meaningful description of the
    swing (e.g. '-140%'). That case must resolve to NaN/N-M even though the
    dollar variance is large enough to be material."""

    def _toy_pnl(self, actual_oi_sign, budget_oi_sign):
        # Revenue/COGS/opex chosen so Operating Income lands on the requested
        # side of zero for actual and budget respectively.
        actual_rev = 100_000 if actual_oi_sign > 0 else 50_000
        budget_rev = 100_000 if budget_oi_sign > 0 else 50_000
        rows = [
            {"month": "2099-03", "line_item": "Revenue", "actual": actual_rev, "budget": budget_rev, "forecast": budget_rev},
            {"month": "2099-03", "line_item": "COGS", "actual": 40_000, "budget": 40_000, "forecast": 40_000},
            {"month": "2099-03", "line_item": "Salaries", "actual": 20_000, "budget": 20_000, "forecast": 20_000},
            {"month": "2099-03", "line_item": "Marketing", "actual": 20_000, "budget": 20_000, "forecast": 20_000},
            {"month": "2099-03", "line_item": "Facilities", "actual": 5_000, "budget": 5_000, "forecast": 5_000},
            {"month": "2099-03", "line_item": "Professional Services", "actual": 3_000, "budget": 3_000, "forecast": 3_000},
            {"month": "2099-03", "line_item": "Other", "actual": 1_000, "budget": 1_000, "forecast": 1_000},
        ]
        return variance.add_subtotals(pd.DataFrame(rows))

    def test_oi_sign_flip_is_not_meaningful(self):
        # Actual OI negative (Revenue 50k - 89k costs = -39k), Budget OI positive (100k - 89k = +11k).
        pnl = self._toy_pnl(actual_oi_sign=-1, budget_oi_sign=1)
        var_df = variance.compute_variance(pnl)
        oi_row = var_df[var_df.line_item == "Operating Income"].iloc[0]
        self.assertTrue(pd.isna(oi_row["variance_pct_vs_budget"]),
                         "a profit-metric sign flip must produce NaN (N/M), not a misleading percentage")
        # The dollar variance and materiality flag must still be intact —
        # only the percentage is suppressed.
        self.assertAlmostEqual(oi_row["variance_vs_budget"], -39_000 - 11_000)
        self.assertTrue(oi_row["flag_vs_budget"], "a $50k swing must still flag as material via the dollar threshold")

    def test_no_sign_flip_still_computes_a_normal_percentage(self):
        # Both actual and budget OI positive -> ordinary percentage, not NaN.
        pnl = self._toy_pnl(actual_oi_sign=1, budget_oi_sign=1)
        var_df = variance.compute_variance(pnl)
        oi_row = var_df[var_df.line_item == "Operating Income"].iloc[0]
        self.assertTrue(pd.notna(oi_row["variance_pct_vs_budget"]))

    def test_real_dataset_flags_the_known_oi_sign_flip_months(self):
        """Regression check: Feb/Jun/Jul 2025 are known (from the real
        synthetic data) to have actual OI on the opposite side of zero from
        budget OI — those months' % variance must be N/M."""
        var_df = _real_pnl_and_variance()
        oi = var_df[var_df.line_item == "Operating Income"].set_index("month")
        for month in ("2025-02", "2025-06", "2025-07"):
            self.assertTrue(pd.isna(oi.loc[month, "variance_pct_vs_budget"]),
                             f"{month}: expected N/M (sign flip), got {oi.loc[month, 'variance_pct_vs_budget']}")
            self.assertTrue(oi.loc[month, "flag_vs_budget"], f"{month}: should still be material via $ threshold")


class TestMaterialityFilteredCommentary(unittest.TestCase):
    """Executive commentary must never let an immaterial variance crowd out
    or dilute the real (material) drivers."""

    def _toy_pnl(self):
        rows = [
            {"line_item": "Revenue", "actual": 100_000, "budget": 90_000, "forecast": 90_000},   # material, favorable
            {"line_item": "COGS", "actual": 40_000, "budget": 40_050, "forecast": 40_050},         # immaterial ($50, 0.1%)
            {"line_item": "Salaries", "actual": 20_000, "budget": 20_000, "forecast": 20_000},     # immaterial ($0)
            {"line_item": "Marketing", "actual": 12_000, "budget": 8_000, "forecast": 8_000},      # material, unfavorable
            {"line_item": "Facilities", "actual": 5_010, "budget": 5_000, "forecast": 5_000},      # immaterial ($10)
            {"line_item": "Professional Services", "actual": 3_000, "budget": 3_000, "forecast": 3_000},  # immaterial
            {"line_item": "Other", "actual": 1_000, "budget": 1_000, "forecast": 1_000},           # immaterial
        ]
        for row in rows:
            row["month"] = "2099-04"
        return variance.add_subtotals(pd.DataFrame(rows))

    def test_material_top_drivers_excludes_immaterial_lines(self):
        var_df = variance.compute_variance(self._toy_pnl())
        favorable, unfavorable = variance.material_top_drivers(var_df, "2099-04", n=5)
        favorable_lines = {line for line, _ in favorable}
        unfavorable_lines = {line for line, _ in unfavorable}

        self.assertIn("Revenue", favorable_lines)
        self.assertIn("Marketing", unfavorable_lines)
        for immaterial_line in ("COGS", "Salaries", "Facilities", "Professional Services", "Other"):
            self.assertNotIn(immaterial_line, favorable_lines | unfavorable_lines,
                              f"{immaterial_line} is immaterial and must not appear in executive commentary")

    def test_narrative_bullets_only_cite_material_drivers(self):
        var_df = variance.compute_variance(self._toy_pnl())
        summary = variance.narrative(var_df, "2099-04", n=5)
        all_text = " ".join(summary["favorable"] + summary["unfavorable"])
        for immaterial_line in ("COGS:", "Salaries:", "Facilities:", "Professional Services:", "Other:"):
            self.assertNotIn(immaterial_line, all_text)


class TestFormatK(unittest.TestCase):
    def test_small_amounts_fall_back_to_whole_dollars(self):
        self.assertEqual(variance.format_k(476), "$476")
        self.assertEqual(variance.format_k(-397), "-$397")

    def test_large_amounts_use_k_suffix(self):
        self.assertEqual(variance.format_k(83_268), "$83K")
        self.assertEqual(variance.format_k(-58_521), "-$59K")


class TestVarianceWorkbookRegressions(unittest.TestCase):
    """Guards two specific regressions: the Summary sheet's company-name
    title going missing or diverging from the single configured source of
    truth, and Variance Detail's freeze panes being overwritten with
    something other than the fixed 'C4' (e.g. a data-dependent row).
    Builds a FRESH workbook in an isolated temp dir — not just inspecting
    output/variance_analysis.xlsx — so this fails if the bug is
    reintroduced by anyone, not only if the committed output goes stale."""

    @classmethod
    def setUpClass(cls):
        import generate_data

        cls.tmp_dir = tempfile.TemporaryDirectory()
        tmp_path = Path(cls.tmp_dir.name)
        db_path = tmp_path / "close_pack.db"
        data = generate_data.generate()
        generate_data._write_sqlite(data, db_path)

        cls.var_df = variance.compute_variance(variance.load_pnl(db_path))
        cls.month = max(cls.var_df["month"])
        cls.workbook_path = tmp_path / "variance_analysis.xlsx"
        excel_export.build_variance_workbook(cls.var_df, cls.month, cls.workbook_path)

        import openpyxl
        cls.wb = openpyxl.load_workbook(cls.workbook_path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp_dir.cleanup()

    def test_summary_a1_equals_configured_company_name(self):
        """Summary!A1 must equal excel_export.COMPANY_SHORT_NAME exactly —
        proves A1 is DERIVED from the one configured company name, not an
        independently hardcoded literal that could drift from it."""
        ws = self.wb["Summary"]
        self.assertEqual(ws.cell(row=1, column=1).value, excel_export.COMPANY_SHORT_NAME)
        self.assertEqual(excel_export.COMPANY_SHORT_NAME, "Beacon Outdoor Goods")

    def test_variance_detail_freeze_panes_is_exactly_c4(self):
        ws = self.wb["Variance Detail"]
        self.assertEqual(ws.freeze_panes, "C4",
                          "freeze_panes must be exactly 'C4' (rows 1-3 + columns A-B), "
                          "never a data-dependent row like 'C99'")

    def test_pnl_reconciliation_still_holds_in_the_fresh_workbook(self):
        for month, group in self.var_df.groupby("month"):
            by_line = group.set_index("line_item")
            for scenario in ("actual", "budget", "forecast"):
                gp = by_line.loc["Revenue", scenario] - by_line.loc["COGS", scenario]
                opex = by_line.loc[variance.OPEX_LINES, scenario].sum()
                self.assertAlmostEqual(gp - opex, by_line.loc["Operating Income", scenario], places=2,
                                       msg=f"OI reconciliation broke for {month}/{scenario}")

    def test_materiality_logic_still_holds_in_the_fresh_workbook(self):
        cogs = self.var_df[self.var_df.line_item == "COGS"].set_index("month")
        self.assertTrue(cogs.loc["2025-06", "flag_vs_budget"])
        self.assertTrue(cogs.loc["2025-07", "flag_vs_budget"])

    def test_waterfall_reconciliation_still_exact_in_the_fresh_workbook(self):
        for month in sorted(self.var_df["month"].unique()):
            recon = variance.waterfall_reconciliation(self.var_df, month)
            self.assertEqual(recon["difference"], 0.0, f"waterfall did not reconcile for {month}")


if __name__ == "__main__":
    unittest.main()
